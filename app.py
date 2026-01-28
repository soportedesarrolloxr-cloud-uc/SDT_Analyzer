import streamlit as st
import pandas as pd
import io
import json
from datetime import datetime
import os
from openai import OpenAI
import PyPDF2
import docx2txt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import tempfile

# Configuración de la página
st.set_page_config(
    page_title="Sistema de Análisis de Admisión (SDT)",
    page_icon="🎓",
    layout="wide"
)

# Estilos CSS personalizados
st.markdown("""
<style>
    .main {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 20px;
    }
    .stApp {
        background: transparent;
    }
    div[data-testid="stFileUploader"] {
        background: white;
        border-radius: 15px;
        padding: 20px;
        border: 3px dashed #667eea;
    }
    .upload-text {
        text-align: center;
        color: #333;
        font-size: 1.1em;
    }
    .info-card {
        background: white;
        padding: 20px;
        border-radius: 10px;
        box-shadow: 0 2px 10px rgba(0, 0, 0, 0.1);
        margin: 10px 0;
    }
    .score-circle {
        display: inline-flex;
        width: 80px;
        height: 80px;
        border-radius: 50%;
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        align-items: center;
        justify-content: center;
        font-size: 1.8em;
        font-weight: bold;
        margin: 10px;
    }
    h1, h2, h3 {
        color: white !important;
    }
</style>
""", unsafe_allow_html=True)

# Inicializar cliente OpenAI
@st.cache_resource
def get_openai_client():
    # Primero intenta obtener de secrets (Streamlit Cloud)
    api_key = None
    try:
        api_key = st.secrets["OPENAI_API_KEY"]
    except:
        # Si no está en secrets, intenta con variable de entorno (local)
        api_key = os.getenv('OPENAI_API_KEY')
    
    if not api_key:
        st.error("⚠️ No se encontró la clave API de OpenAI")
        st.info("**Para uso local:** Configura la variable de entorno OPENAI_API_KEY en tu archivo .env")
        st.info("**Para Streamlit Cloud:** Configura OPENAI_API_KEY en Settings → Secrets")
        st.stop()
    return OpenAI(api_key=api_key)

client = get_openai_client()

# Función para extraer texto de PDF
def extract_text_from_pdf(file):
    try:
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text
    except Exception as e:
        st.error(f"Error al leer PDF: {str(e)}")
        return None

# Función para extraer texto de DOCX
def extract_text_from_docx(file):
    try:
        text = docx2txt.process(file)
        return text
    except Exception as e:
        st.error(f"Error al leer DOCX: {str(e)}")
        return None

# Función para extraer texto de TXT
def extract_text_from_txt(file):
    try:
        return file.read().decode('utf-8')
    except Exception as e:
        st.error(f"Error al leer TXT: {str(e)}")
        return None

# Función para leer Excel/CSV
def read_excel_file(file):
    try:
        file_extension = file.name.split('.')[-1].lower()
        if file_extension == 'csv':
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
        return df
    except Exception as e:
        st.error(f"Error al leer archivo Excel/CSV: {str(e)}")
        return None

# Función principal de análisis con OpenAI
def analyze_admission_form(text_content):
    system_prompt = """ROL Y CONTEXTO:
Actúa como experto en Psicología Educativa y Motivación, especializado en la Teoría de la Autodeterminación (Self-Determination Theory, SDT) de Ryan y Deci.

MARCO TEÓRICO - TEORÍA DE LA AUTODETERMINACIÓN (SDT):
Evalúa la motivación académica basándote en:
- Motivación intrínseca: interés genuino, disfrute, curiosidad
- Regulaciones extrínsecas:
  * Externa: recompensas/presiones externas
  * Introyectada: presión interna, culpa, orgullo
  * Identificada: utilidad para metas personales
  * Integrada: coherencia con identidad y valores
- Amotivación: sin razón clara o desinterés
- Necesidades psicológicas básicas: autonomía, competencia, relación

CONTEXTO INSTITUCIONAL:
Universidad Continental - Modalidad a Distancia
Población diversa: todo Perú, 18+ años, muchos trabajan y tienen familia
Objetivo: diagnosticar niveles motivacionales, NO filtrar

CRITERIOS DE EVALUACIÓN (Escala 1-6):
Cada respuesta se evalúa según nivel motivacional:
6 = Motivación Intrínseca (interés genuino, disfrute)
5 = Motivación Integrada (coherencia con identidad y valores)
4 = Motivación Identificada (utilidad personal significativa)
3 = Motivación Introyectada (presión interna, orgullo, culpa)
2 = Motivación Extrínseca (recompensas/presiones externas)
1 = Amotivación (sin razón clara, desinterés)

ESTRUCTURA DE ANÁLISIS:
Debes analizar el formulario buscando evidencia de estos niveles motivacionales y responder SOLO con un objeto JSON válido (sin markdown, sin ```json) con esta estructura:

{
  "informacion_extraida": {
    "nombre": "...",
    "edad": "...",
    "programa": "...",
    "otros_campos": {}
  },
  "evaluacion_motivacional": {
    "eleccion_carrera": {
      "puntaje": 1-6,
      "tipo_motivacion": "Intrínseca/Integrada/Identificada/Introyectada/Extrínseca/Amotivación",
      "justificacion": "Evidencia textual breve (1-2 líneas)"
    },
    "experiencia_relacionada": {
      "puntaje": 1-6,
      "tipo_motivacion": "...",
      "justificacion": "..."
    },
    "uso_futuro": {
      "puntaje": 1-6,
      "tipo_motivacion": "...",
      "justificacion": "..."
    }
  },
  "necesidades_psicologicas": {
    "autonomia": "Alta/Media/Baja - Breve análisis",
    "competencia": "Alta/Media/Baja - Breve análisis",
    "relacion": "Alta/Media/Baja - Breve análisis"
  },
  "calificacion_real": "suma de puntajes (máx 18)",
  "calificacion_sobre_20": "calificación_real/18 * 20 (dos decimales)",
  "recomendaciones": "Basadas en SDT: fortalecer autonomía, competencia o relación según necesidad",
  "nivel_motivacional_general": "Predominantemente Intrínseco/Integrado/Identificado/Introyectado/Extrínseco/Amotivado"
}

IMPORTANTE:
- Responde ÚNICAMENTE con el JSON (sin markdown ni bloques de código)
- Si hay campos obligatorios faltantes, inferir lo posible del texto
- Calcula calificacion_sobre_20 = (calificacion_real / 18) * 20
- Redondea a 2 decimales
"""

    user_prompt = f"""Analiza el siguiente formulario de admisión y proporciona un análisis completo según SDT:

{text_content}

Recuerda: responde ÚNICAMENTE con el JSON, sin ningún texto adicional ni bloques de código markdown."""

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt}
            ],
            temperature=0.3,
            max_tokens=2000
        )
        
        content = response.choices[0].message.content.strip()
        
        # Limpiar respuesta de posibles markdown
        if content.startswith('```'):
            content = content.split('```')[1]
            if content.startswith('json'):
                content = content[4:]
            content = content.strip()
        
        result = json.loads(content)
        result['tokens_used'] = response.usage.total_tokens
        result['timestamp'] = datetime.now().isoformat()
        result['success'] = True
        
        return result
        
    except json.JSONDecodeError as e:
        st.error(f"Error al parsear respuesta JSON: {str(e)}")
        return {"success": False, "error": "Error al parsear respuesta de IA"}
    except Exception as e:
        st.error(f"Error en análisis con OpenAI: {str(e)}")
        return {"success": False, "error": str(e)}

# Función para procesar registros de Excel
def process_excel_records(df, progress_bar, status_text):
    results = []
    total = len(df)
    
    required_fields = ['Respuesta 1', 'Respuesta 2', 'Respuesta 3']
    
    for idx, row in df.iterrows():
        status_text.text(f"Procesando registro {idx + 1} de {total}...")
        progress_bar.progress((idx + 1) / total)
        
        # Construir texto del formulario
        form_text = f"""
Nombre: {row.get('Nombre', 'N/A')}
Apellidos: {row.get('Apellidos', 'N/A')}
Correo: {row.get('Correo electrónico', 'N/A')}
Edad: {row.get('Edad', 'N/A')}
Programa: {row.get('Programa', 'N/A')}

Pregunta 1: ¿Por qué elegiste esta carrera?
Respuesta 1: {row.get('Respuesta 1', 'Sin respuesta')}

Pregunta 2: ¿Qué experiencia tienes relacionada con esta carrera?
Respuesta 2: {row.get('Respuesta 2', 'Sin respuesta')}

Pregunta 3: ¿Cómo planeas usar lo que aprendas?
Respuesta 3: {row.get('Respuesta 3', 'Sin respuesta')}
"""
        
        # Validar que tenga al menos las respuestas requeridas
        missing_fields = [field for field in required_fields if pd.isna(row.get(field)) or str(row.get(field)).strip() == '']
        
        if missing_fields:
            results.append({
                'success': False,
                'registro_numero': idx + 1,
                'nombre': row.get('Nombre', 'N/A'),
                'apellidos': row.get('Apellidos', 'N/A'),
                'correo': row.get('Correo electrónico', 'N/A'),
                'error': f"Campos faltantes: {', '.join(missing_fields)}"
            })
            continue
        
        # Analizar con OpenAI
        analysis = analyze_admission_form(form_text)
        
        result = {
            'registro_numero': idx + 1,
            'nombre': row.get('Nombre', 'N/A'),
            'apellidos': row.get('Apellidos', 'N/A'),
            'correo': row.get('Correo electrónico', 'N/A'),
            'success': analysis.get('success', False),
        }
        
        if analysis.get('success'):
            result['analysis'] = analysis
        else:
            result['error'] = analysis.get('error', 'Error desconocido')
        
        results.append(result)
    
    return results

# Función para generar Excel de resultados
def generate_excel_report(results):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados Análisis SDT"
    
    # Estilos
    header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Encabezados
    headers = [
        'N°', 'Nombre', 'Apellidos', 'Correo', 'Calif. Real', 'Calif. /20',
        'R1 Punt.', 'R1 Justificación', 'R1 Tipo',
        'R2 Punt.', 'R2 Justificación', 'R2 Tipo',
        'R3 Punt.', 'R3 Justificación', 'R3 Tipo',
        'Nivel General', 'Autonomía', 'Competencia', 'Relación', 'Recomendaciones'
    ]
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    
    # Ajustar anchos de columna
    column_widths = [5, 15, 15, 25, 10, 10, 8, 40, 15, 8, 40, 15, 8, 40, 15, 20, 30, 30, 30, 50]
    for col_num, width in enumerate(column_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_num).column_letter].width = width
    
    # Datos
    for result in results:
        r = result
        a = r.get('analysis', {})
        
        ws.append([
            r.get('registro_numero', ''),
            r.get('nombre', ''),
            r.get('apellidos', ''),
            r.get('correo', ''),
            a.get('calificacion_real', ''),
            a.get('calificacion_sobre_20', ''),
            a.get('evaluacion_motivacional', {}).get('eleccion_carrera', {}).get('puntaje', ''),
            a.get('evaluacion_motivacional', {}).get('eleccion_carrera', {}).get('justificacion', ''),
            a.get('evaluacion_motivacional', {}).get('eleccion_carrera', {}).get('tipo_motivacion', ''),
            a.get('evaluacion_motivacional', {}).get('experiencia_relacionada', {}).get('puntaje', ''),
            a.get('evaluacion_motivacional', {}).get('experiencia_relacionada', {}).get('justificacion', ''),
            a.get('evaluacion_motivacional', {}).get('experiencia_relacionada', {}).get('tipo_motivacion', ''),
            a.get('evaluacion_motivacional', {}).get('uso_futuro', {}).get('puntaje', ''),
            a.get('evaluacion_motivacional', {}).get('uso_futuro', {}).get('justificacion', ''),
            a.get('evaluacion_motivacional', {}).get('uso_futuro', {}).get('tipo_motivacion', ''),
            a.get('nivel_motivacional_general', ''),
            a.get('necesidades_psicologicas', {}).get('autonomia', ''),
            a.get('necesidades_psicologicas', {}).get('competencia', ''),
            a.get('necesidades_psicologicas', {}).get('relacion', ''),
            a.get('recomendaciones', '')
        ])
    
    # Ajustar altura de filas
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        ws.row_dimensions[row[0].row].height = 30
        for cell in row:
            cell.alignment = Alignment(vertical='middle', wrap_text=True)
            cell.border = border
    
    # Guardar en buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

# INTERFAZ PRINCIPAL
def main():
    st.title("🎓 Sistema de Análisis de Admisión (SDT)")
    st.markdown("### Sube formularios individuales (PDF/DOCX/TXT) o múltiples registros (Excel/CSV)")
    
    # Verificar API Key
    if not os.getenv('OPENAI_API_KEY'):
        st.error("⚠️ No se encontró la clave API de OpenAI")
        st.info("Por favor, configura la variable de entorno OPENAI_API_KEY antes de usar la aplicación")
        st.stop()
    
    # Upload de archivo
    uploaded_file = st.file_uploader(
        "📎 Selecciona un archivo",
        type=['pdf', 'docx', 'doc', 'txt', 'xlsx', 'xls', 'csv'],
        help="Formatos soportados: PDF, DOCX, TXT para análisis individual | XLSX, XLS, CSV para análisis masivo"
    )
    
    if uploaded_file:
        file_extension = uploaded_file.name.split('.')[-1].lower()
        file_size = uploaded_file.size / (1024 * 1024)  # MB
        
        st.info(f"📄 **Archivo:** {uploaded_file.name} ({file_size:.2f} MB)")
        
        is_batch = file_extension in ['xlsx', 'xls', 'csv']
        
        if is_batch:
            st.warning("📊 **Modo Masivo Detectado** - Se procesarán múltiples registros")
        
        # Botón de análisis
        if st.button("🚀 Analizar Formulario(s)", type="primary", use_container_width=True):
            with st.spinner("Procesando..."):
                if is_batch:
                    # PROCESAMIENTO MASIVO
                    df = read_excel_file(uploaded_file)
                    
                    if df is not None:
                        st.success(f"✅ {len(df)} registros encontrados")
                        
                        progress_bar = st.progress(0)
                        status_text = st.empty()
                        
                        results = process_excel_records(df, progress_bar, status_text)
                        
                        status_text.text("✅ Análisis completado!")
                        progress_bar.progress(1.0)
                        
                        # Guardar en session state
                        st.session_state['batch_results'] = results
                        st.session_state['batch_filename'] = uploaded_file.name
                        
                        # Mostrar resultados
                        st.markdown("---")
                        st.header("📊 Resultados del Análisis Masivo")
                        
                        success_count = sum(1 for r in results if r.get('success'))
                        avg_score = sum(
                            float(r['analysis']['calificacion_sobre_20']) 
                            for r in results 
                            if r.get('success') and r.get('analysis', {}).get('calificacion_sobre_20')
                        ) / success_count if success_count > 0 else 0
                        
                        col1, col2, col3, col4 = st.columns(4)
                        with col1:
                            st.metric("📋 Total Registros", len(results))
                        with col2:
                            st.metric("✅ Procesados", success_count)
                        with col3:
                            st.metric("❌ Errores", len(results) - success_count)
                        with col4:
                            st.metric("📊 Promedio", f"{avg_score:.2f}/20")
                        
                        # Botón de descarga
                        excel_buffer = generate_excel_report(results)
                        st.download_button(
                            label="📥 Descargar Resultados (Excel)",
                            data=excel_buffer,
                            file_name=f"Resultados_Analisis_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            type="primary"
                        )
                        
                        # Tabla de resultados
                        st.markdown("### Detalle por Postulante")
                        for i, result in enumerate(results):
                            with st.expander(f"**{result['registro_numero']}. {result['apellidos']}, {result['nombre']}** - {result['correo']}"):
                                if result.get('success'):
                                    analysis = result['analysis']
                                    
                                    col1, col2 = st.columns([1, 3])
                                    with col1:
                                        st.markdown(f"<div style='text-align: center;'><div class='score-circle'>{analysis.get('calificacion_sobre_20', 'N/A')}</div></div>", unsafe_allow_html=True)
                                    with col2:
                                        st.markdown(f"**Calificación Real:** {analysis.get('calificacion_real', 'N/A')}/18")
                                        st.markdown(f"**Nivel Motivacional:** {analysis.get('nivel_motivacional_general', 'N/A')}")
                                    
                                    # Evaluación motivacional
                                    st.markdown("#### 📝 Evaluación Motivacional")
                                    eval_mot = analysis.get('evaluacion_motivacional', {})
                                    
                                    col1, col2, col3 = st.columns(3)
                                    with col1:
                                        if 'eleccion_carrera' in eval_mot:
                                            st.info(f"**R1:** {eval_mot['eleccion_carrera'].get('puntaje')}/6\n\n{eval_mot['eleccion_carrera'].get('tipo_motivacion')}")
                                    with col2:
                                        if 'experiencia_relacionada' in eval_mot:
                                            st.info(f"**R2:** {eval_mot['experiencia_relacionada'].get('puntaje')}/6\n\n{eval_mot['experiencia_relacionada'].get('tipo_motivacion')}")
                                    with col3:
                                        if 'uso_futuro' in eval_mot:
                                            st.info(f"**R3:** {eval_mot['uso_futuro'].get('puntaje')}/6\n\n{eval_mot['uso_futuro'].get('tipo_motivacion')}")
                                    
                                    # Necesidades psicológicas
                                    if 'necesidades_psicologicas' in analysis:
                                        st.markdown("#### 🧠 Necesidades Psicológicas (SDT)")
                                        nec = analysis['necesidades_psicologicas']
                                        col1, col2, col3 = st.columns(3)
                                        with col1:
                                            st.success(f"**Autonomía:**\n{nec.get('autonomia', 'N/A')}")
                                        with col2:
                                            st.success(f"**Competencia:**\n{nec.get('competencia', 'N/A')}")
                                        with col3:
                                            st.success(f"**Relación:**\n{nec.get('relacion', 'N/A')}")
                                    
                                    # Recomendaciones
                                    if 'recomendaciones' in analysis:
                                        st.markdown("#### 💡 Recomendaciones")
                                        st.info(analysis['recomendaciones'])
                                else:
                                    st.error(f"❌ Error: {result.get('error', 'Error desconocido')}")
                
                else:
                    # PROCESAMIENTO INDIVIDUAL
                    text_content = None
                    
                    if file_extension == 'pdf':
                        text_content = extract_text_from_pdf(uploaded_file)
                    elif file_extension in ['docx', 'doc']:
                        text_content = extract_text_from_docx(uploaded_file)
                    elif file_extension == 'txt':
                        text_content = extract_text_from_txt(uploaded_file)
                    
                    if text_content and text_content.strip():
                        st.success(f"✅ Texto extraído ({len(text_content)} caracteres)")
                        
                        analysis = analyze_admission_form(text_content)
                        
                        if analysis.get('success'):
                            st.markdown("---")
                            st.header("📊 Resultado del Análisis")
                            
                            # Información básica
                            info = analysis.get('informacion_extraida', {})
                            col1, col2, col3 = st.columns(3)
                            with col1:
                                st.info(f"**Nombre:** {info.get('nombre', 'N/A')}")
                            with col2:
                                st.info(f"**Edad:** {info.get('edad', 'N/A')}")
                            with col3:
                                st.info(f"**Programa:** {info.get('programa', 'N/A')}")
                            
                            # Calificaciones
                            col1, col2 = st.columns([1, 2])
                            with col1:
                                st.markdown(f"<div style='text-align: center;'><div class='score-circle'>{analysis.get('calificacion_sobre_20', 'N/A')}</div><p style='text-align: center; color: white;'>Calificación sobre 20</p></div>", unsafe_allow_html=True)
                            with col2:
                                st.metric("Calificación Real", f"{analysis.get('calificacion_real', 'N/A')}/18")
                                st.metric("Nivel Motivacional", analysis.get('nivel_motivacional_general', 'N/A'))
                            
                            # Evaluación motivacional
                            st.markdown("### 📝 Evaluación Motivacional Detallada")
                            eval_mot = analysis.get('evaluacion_motivacional', {})
                            
                            for key, label in [
                                ('eleccion_carrera', 'R1: Elección de Carrera'),
                                ('experiencia_relacionada', 'R2: Experiencia Relacionada'),
                                ('uso_futuro', 'R3: Uso Futuro')
                            ]:
                                if key in eval_mot:
                                    item = eval_mot[key]
                                    with st.expander(f"**{label}** - {item.get('puntaje')}/6 ({item.get('tipo_motivacion')})"):
                                        st.write(f"**Justificación:** {item.get('justificacion')}")
                            
                            # Necesidades psicológicas
                            if 'necesidades_psicologicas' in analysis:
                                st.markdown("### 🧠 Necesidades Psicológicas (SDT)")
                                nec = analysis['necesidades_psicologicas']
                                col1, col2, col3 = st.columns(3)
                                with col1:
                                    st.success(f"**Autonomía:**\n\n{nec.get('autonomia', 'N/A')}")
                                with col2:
                                    st.success(f"**Competencia:**\n\n{nec.get('competencia', 'N/A')}")
                                with col3:
                                    st.success(f"**Relación:**\n\n{nec.get('relacion', 'N/A')}")
                            
                            # Recomendaciones
                            if 'recomendaciones' in analysis:
                                st.markdown("### 💡 Recomendaciones")
                                st.info(analysis['recomendaciones'])
                            
                            # Metadata
                            st.markdown("---")
                            st.caption(f"**Archivo:** {uploaded_file.name} | **Tokens:** {analysis.get('tokens_used', 'N/A')} | **Procesado:** {datetime.fromisoformat(analysis.get('timestamp')).strftime('%Y-%m-%d %H:%M:%S')}")
                        
                        else:
                            st.error(f"❌ Error: {analysis.get('error', 'Error desconocido')}")
                    else:
                        st.error("❌ No se pudo extraer texto del archivo o el archivo está vacío")

if __name__ == "__main__":
    main()
